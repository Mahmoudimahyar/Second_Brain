"""Generate a tiny, entirely FICTIONAL dataset for the SecBrain quickstart.

Two sources with different trust tiers:

* ``reference_2024-25.xlsx`` — an "official" spreadsheet (tier L1): five made-up schools with a
  few published facts each. This is the ground truth the engine will never overwrite.
* ``demo_forum.jsonl`` / ``demo_forum_comments.jsonl`` — a made-up community forum (tier L5) in Reddit's
  export shape. People refer to the schools by nicknames, ask questions, and sometimes
  get the facts wrong — which is the whole point.

Every name, number and sentence here is invented. No real institution, person or post.

Run:  python examples/quickstart/make_demo_data.py [output_dir]
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook

SCHOOLS = [
    # name,                                        state, city,        tuition, avg_score
    ("Northfield University School of Dentistry", "MN", "Northfield", 61500, 21.4),
    ("Lakemont College of Dental Medicine",       "NY", "Lakemont",   88200, 22.1),
    ("Ashgrove University School of Dentistry",   "OR", "Ashgrove",   54300, 20.6),
    ("Pinecrest School of Dental Medicine",       "CO", "Pinecrest",  73900, 21.0),
    ("Marlowe University College of Dentistry",   "TX", "Marlowe",    47800, 20.2),
]

# created_utc values are spread over several years so temporal logic has something to chew on
_T2019, _T2022, _T2024, _T2025 = 1560000000, 1655000000, 1718000000, 1745000000

POSTS = [
    ("p01", "molar_mike",    _T2025, "Acceptances",  "Got into Northfield University School of Dentistry!",
     "Still shaking. Four years of work. Anyone else heading to Northfield this fall?"),
    ("p02", "flossophy",     _T2025, "Cost",         "Is Lakemont College of Dental Medicine worth the price?",
     "Tuition at Lakemont looks brutal compared with everywhere else I applied. Is the clinic time worth it?"),
    ("p03", "enamel_annie",  _T2019, "Cost",         "Marlowe tuition is about 30k a year",
     "My cousin went to Marlowe University College of Dentistry and said tuition was around 30k. Cheapest by far."),
    ("p04", "drill_sgt",     _T2024, "Interviews",   "Ashgrove interview — what they asked me",
     "Ashgrove University School of Dentistry asked: Why dentistry and not medicine? "
     "Tell us about a time you failed. How do you handle a patient who refuses treatment?"),
    ("p05", "gap_year_gary", _T2024, "Advice",       "Pinecrest vs Ashgrove for someone who wants early clinic exposure",
     "Accepted at Pinecrest School of Dental Medicine and Ashgrove. Pinecrest starts clinic in year two."),
    ("p06", "molar_mike",    _T2022, "Cost",         "Lakemont raised tuition again",
     "Lakemont College of Dental Medicine went up again this cycle. Loans are going to be rough."),
    ("p07", "plaque_attack", _T2025, "Advice",       "How many shadowing hours is enough?",
     "I have about 80 hours across two offices. Should I keep going or focus on my entrance exam?"),
    ("p08", "flossophy",     _T2024, "Interviews",   "Northfield interview was relaxed",
     "Northfield asked mostly conversational questions. What do you do for fun? Why Minnesota?"),
    ("p09", "crown_jules",   _T2025, "Cost",         "Marlowe is NOT 30k anymore",
     "People keep repeating that Marlowe University College of Dentistry costs 30k. It has been close to "
     "48k for a while now. Check the official numbers before you plan your loans."),
    ("p10", "enamel_annie",  _T2025, "Acceptances",  "Waitlisted at Pinecrest",
     "Pinecrest School of Dental Medicine put me on the waitlist. Did anyone get off it last year?"),
]

COMMENTS = [
    ("c01", "p02", "t3_p02", "crown_jules",   _T2025, "Lakemont is expensive but the patient volume is unreal. Worth it for me."),
    ("c02", "p02", "t3_p02", "drill_sgt",     _T2025, "Not worth it. You can get the same license for half the debt somewhere else."),
    ("c03", "p02", "t1_c02", "flossophy",     _T2025, "That is what I am afraid of. The debt number keeps me up at night."),
    ("c04", "p03", "t3_p03", "gap_year_gary", _T2019, "Can confirm, Marlowe was around 30k back then."),
    ("c05", "p09", "t3_p09", "molar_mike",    _T2025, "Thank you. I almost built my whole budget on the 30k figure."),
    ("c06", "p09", "t3_p09", "plaque_attack", _T2025, "Still the cheapest option I found, even at the higher number."),
    ("c07", "p04", "t3_p04", "enamel_annie",  _T2024, "They asked me the failure question too. Have a real story ready."),
    ("c08", "p05", "t3_p05", "crown_jules",   _T2024, "Early clinic at Pinecrest was the deciding factor for me. No regrets."),
    ("c09", "p01", "t3_p01", "flossophy",     _T2025, "Congrats! See you at Northfield orientation."),
    ("c10", "p10", "t3_p10", "drill_sgt",     _T2025, "Two people in my class came off the Pinecrest waitlist in May."),
    ("c11", "p07", "t3_p07", "gap_year_gary", _T2025, "80 is plenty. Put the time into the exam."),
    ("c12", "p06", "t3_p06", "crown_jules",   _T2022, "Every school went up that year, to be fair."),
]


def write_reference(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tab1"
    ws["A1"] = "Table 1: Published school facts (fictional demo data)"
    ws["A2"] = "Return to Table of Contents"
    for col, header in zip("ABCDE", ("School", "State", "City", "Tuition Resident", "Avg Entrance Score"),
                           strict=True):
        ws[f"{col}5"] = header
    for row, school in enumerate(SCHOOLS, start=6):
        for col, value in zip("ABCDE", school, strict=True):
            ws[f"{col}{row}"] = value
    wb.save(path)


def write_forum(posts_path: Path, comments_path: Path) -> None:
    with posts_path.open("w", encoding="utf-8", newline="\n") as fh:
        for pid, author, created, flair, title, body in POSTS:
            fh.write(json.dumps({
                "id": pid, "author": author, "created_utc": created, "title": title,
                "selftext": body, "link_flair_text": flair, "score": 12, "ups": 12, "downs": 0,
                "num_comments": sum(1 for c in COMMENTS if c[1] == pid),
                "permalink": f"/r/demo_forum/comments/{pid}/", "subreddit": "demo_forum",
            }) + "\n")
    with comments_path.open("w", encoding="utf-8", newline="\n") as fh:
        for cid, pid, parent, author, created, body in COMMENTS:
            fh.write(json.dumps({
                "id": cid, "author": author, "created_utc": created, "body": body,
                "link_id": f"t3_{pid}", "parent_id": parent, "score": 4, "ups": 4, "downs": 0,
                "controversiality": 0, "subreddit": "demo_forum",
            }) + "\n")


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    out = Path(args[0]) if args else Path(__file__).parent / "demo_data"
    out.mkdir(parents=True, exist_ok=True)
    write_reference(out / "reference_2024-25.xlsx")   # the L1 sheet adapter reads the cycle year from the name
    write_forum(out / "demo_forum.jsonl", out / "demo_forum_comments.jsonl")
    print(f"wrote {len(SCHOOLS)} reference rows, {len(POSTS)} posts, {len(COMMENTS)} comments -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
