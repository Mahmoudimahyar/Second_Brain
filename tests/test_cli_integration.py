"""Integration tests for `src.cli` via typer's `CliRunner`.

Exercises each subcommand against a fresh tmp data dir. The Pass 4 / Pass 3 /
gateway-dependent commands use `MockProvider` indirectly through default_gateway()
returning an empty gateway when no API keys are set, so we monkeypatch keys
or rely on the deterministic ingest/query paths.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from src.cli import app
from src.hitl import HITLQueue


@pytest.fixture
def runner() -> CliRunner:
    return CliRunner()


def _synthetic_adea(tmp_path: Path) -> Path:
    """Tiny ADEA-shape workbook for fast CLI smoke."""

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tab1"
    ws["A1"] = "Table 1: Synthetic Tuition"
    ws["A2"] = "Return to Table of Contents"
    ws["A5"] = "School"
    ws["B5"] = "State"
    ws["C5"] = "City"
    ws["D5"] = "Tuition Resident"
    ws["E5"] = "Avg DAT AA"
    ws["A6"] = "Harvard School of Dental Medicine"
    ws["B6"] = "MA"
    ws["C6"] = "Boston"
    ws["D6"] = 79000
    ws["E6"] = 22.5
    out = tmp_path / "SDE2_2024-25.xlsx"
    wb.save(out)
    return out


def test_stats_on_empty_data_dir(runner: CliRunner, tmp_path: Path) -> None:
    result = runner.invoke(app, ["--data-dir", str(tmp_path), "stats"])
    assert result.exit_code == 0
    assert "Total nodes" in result.output


def test_ingest_l1_adea_smoke(runner: CliRunner, tmp_path: Path) -> None:
    payload = _synthetic_adea(tmp_path)
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path / "data"),
               "ingest", "l1-adea", str(payload)],
    )
    assert result.exit_code == 0
    assert "schools=1" in result.output


def test_stats_after_l1_ingest(runner: CliRunner, tmp_path: Path) -> None:
    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])

    result = runner.invoke(app, ["--data-dir", str(data_dir), "stats"])
    assert result.exit_code == 0
    assert "School" in result.output
    assert "Total nodes" in result.output


def test_query_returns_table(runner: CliRunner, tmp_path: Path) -> None:
    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])

    result = runner.invoke(
        app, ["--data-dir", str(data_dir), "query", "Harvard", "--limit", "5"],
    )
    assert result.exit_code == 0
    assert "query:" in result.output.lower() or "results" in result.output.lower()


def test_canonical_returns_match(runner: CliRunner, tmp_path: Path) -> None:
    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])

    result = runner.invoke(
        app, ["--data-dir", str(data_dir),
               "canonical", "Harvard School of Dental Medicine"],
    )
    assert result.exit_code == 0
    assert "canonical_id" in result.output or "harvard" in result.output.lower()


def test_canonical_returns_no_match_for_unknown(runner: CliRunner, tmp_path: Path) -> None:
    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])

    result = runner.invoke(
        app, ["--data-dir", str(data_dir), "canonical", "Totally Unrelated University"],
    )
    assert result.exit_code == 0
    assert "No L1 match" in result.output


def test_reconcile_demo_runs(runner: CliRunner, tmp_path: Path) -> None:
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "reconcile-demo"],
    )
    assert result.exit_code == 0
    assert "l1_clash_invalidated" in result.output


def test_hitl_list_empty(runner: CliRunner, tmp_path: Path) -> None:
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "list"],
    )
    assert result.exit_code == 0
    assert "HITL queue" in result.output


def test_hitl_pull_returns_none_when_empty(runner: CliRunner, tmp_path: Path) -> None:
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "pull"],
    )
    assert result.exit_code == 0
    assert "No pending items" in result.output


def _enqueue_synthetic_hitl(data_dir: Path) -> str:
    queue = HITLQueue(sqlite_path=data_dir / "sqlite" / "store.db")
    return queue.enqueue(
        item_type="alias_match",
        payload={"alias": "UPenn", "candidates": ["school:upenn"]},
    )


def test_hitl_show_displays_item(runner: CliRunner, tmp_path: Path) -> None:
    item_id = _enqueue_synthetic_hitl(tmp_path)
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "show", item_id],
    )
    assert result.exit_code == 0
    assert item_id in result.output
    assert "alias_match" in result.output
    assert "UPenn" in result.output


def test_hitl_show_reports_missing_item(runner: CliRunner, tmp_path: Path) -> None:
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "show", "hitl:nope"],
    )
    assert result.exit_code != 0


def test_hitl_list_status_filter(runner: CliRunner, tmp_path: Path) -> None:
    _enqueue_synthetic_hitl(tmp_path)
    pending = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "list", "--status", "pending"],
    )
    committed = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "list", "--status", "committed"],
    )
    assert pending.exit_code == 0
    assert committed.exit_code == 0
    assert "alias_match" in pending.output
    assert "alias_match" not in committed.output


def test_hitl_escalate_marks_item_escalated(runner: CliRunner, tmp_path: Path) -> None:
    item_id = _enqueue_synthetic_hitl(tmp_path)
    runner.invoke(app, ["--data-dir", str(tmp_path), "hitl", "pull"])
    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "escalate", item_id,
              "--notes", "needs Mahyar"],
    )
    assert result.exit_code == 0
    show = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "show", item_id],
    )
    assert "escalated" in show.output.lower()


def test_hitl_commit_from_reads_pull_artifact(runner: CliRunner, tmp_path: Path) -> None:
    """`hitl pull` writes a JSON; user edits decision; `commit-from` reads it back."""

    item_id = _enqueue_synthetic_hitl(tmp_path)
    runner.invoke(app, ["--data-dir", str(tmp_path), "hitl", "pull"])
    artifact = tmp_path / "hitl" / f"{item_id}.json"
    assert artifact.is_file()
    payload = json.loads(artifact.read_text(encoding="utf-8"))
    payload["decision"] = {"verdict": "reject", "notes": "bad match"}
    artifact.write_text(json.dumps(payload), encoding="utf-8")

    result = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "commit-from", str(artifact)],
    )
    assert result.exit_code == 0
    show = runner.invoke(
        app, ["--data-dir", str(tmp_path), "hitl", "show", item_id],
    )
    assert "reject" in show.output.lower()
    assert "committed" in show.output.lower()


def test_extract_no_posts_yields_warning(runner: CliRunner, tmp_path: Path) -> None:
    """When the graph has no posts, extract should yield a helpful message."""

    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])
    result = runner.invoke(
        app, ["--data-dir", str(data_dir), "extract",
               "--task", "sentiment", "--sample", "5"],
    )
    assert result.exit_code == 0
    assert "No Posts found" in result.output


def test_cluster_no_posts_yields_warning(runner: CliRunner, tmp_path: Path) -> None:
    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])
    result = runner.invoke(
        app, ["--data-dir", str(data_dir), "cluster",
               "--sample", "5", "--embedder", "hash"],
    )
    assert result.exit_code == 0
    assert "No posts found" in result.output


def test_unknown_extract_task_exits_with_error(runner: CliRunner, tmp_path: Path) -> None:
    """An unknown task name should exit non-zero."""

    payload = _synthetic_adea(tmp_path)
    data_dir = tmp_path / "data"
    runner.invoke(app, ["--data-dir", str(data_dir), "ingest", "l1-adea", str(payload)])
    result = runner.invoke(
        app, ["--data-dir", str(data_dir), "extract", "--task", "bogus"],
    )
    assert result.exit_code != 0
